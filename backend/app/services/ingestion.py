from io import BytesIO
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader
from docx import Document as Docx
from pptx import Presentation
from openpyxl import load_workbook


def load_pdf(b):
    return '\n\n'.join(
        f'[PAGE {i}] {p.extract_text() or ""}'
        for i, p in enumerate(PdfReader(BytesIO(b)).pages, 1)
    )


def load_docx(b):
    return '\n'.join(
        p.text
        for p in Docx(BytesIO(b)).paragraphs
        if p.text.strip()
    )


def load_pptx(b):
    r = []

    for i, s in enumerate(Presentation(BytesIO(b)).slides, 1):
        r.append(
            f'[SLIDE {i}]\n'
            + '\n'.join(
                x.text
                for x in s.shapes
                if hasattr(x, 'text')
            )
        )

    return '\n\n'.join(r)


def load_xlsx(b):
    """
    Convert Excel worksheets into retrieval-friendly text.

    Each row is represented using the column headers so that
    the relationship between a value and its column is preserved.
    """

    workbook = load_workbook(
        BytesIO(b),
        read_only=True,
        data_only=True
    )

    output = []

    for sheet in workbook.worksheets:
        rows = list(
            sheet.iter_rows(values_only=True)
        )

        if not rows:
            continue

        # First non-empty row is treated as the header.
        header_index = None

        for i, row in enumerate(rows):
            if any(
                value is not None and str(value).strip()
                for value in row
            ):
                header_index = i
                break

        if header_index is None:
            continue

        headers = [
            str(value).strip()
            if value is not None and str(value).strip()
            else f'Column {i + 1}'
            for i, value in enumerate(rows[header_index])
        ]

        output.append(
            f'[SHEET {sheet.title}]'
        )

        output.append(
            'COLUMNS: ' + ' | '.join(headers)
        )

        # Convert every data row into explicit key/value pairs.
        for row_number, row in enumerate(
            rows[header_index + 1:],
            header_index + 2
        ):
            values = list(row)

            # Skip completely empty rows.
            if not any(
                value is not None and str(value).strip()
                for value in values
            ):
                continue

            fields = []

            for i, header in enumerate(headers):
                value = (
                    values[i]
                    if i < len(values)
                    else None
                )

                if value is not None and str(value).strip():
                    fields.append(
                        f'{header}: {value}'
                    )

            if fields:
                output.append(
                    f'ROW {row_number}: '
                    + ' | '.join(fields)
                )

        output.append('')

    return '\n'.join(output)


def load_url(url):
    r = requests.get(
        url,
        timeout=20,
        headers={
            'User-Agent':
                'EnterpriseAIKnowledgeAssistant'
        }
    )

    r.raise_for_status()

    s = BeautifulSoup(
        r.text,
        'html.parser'
    )

    for t in s([
        'script',
        'style',
        'noscript'
    ]):
        t.decompose()

    return s.get_text(
        '\n',
        strip=True
    )


def load_file(name, b):
    e = Path(name).suffix.lower()

    if e == '.pdf':
        return load_pdf(b)

    if e == '.docx':
        return load_docx(b)

    if e == '.pptx':
        return load_pptx(b)

    if e == '.xlsx':
        return load_xlsx(b)

    if e in {'.txt', '.md', '.csv'}:
        return b.decode(
            'utf-8',
            'ignore'
        )

    raise ValueError(
        f'Unsupported file type: {e}'
    )
