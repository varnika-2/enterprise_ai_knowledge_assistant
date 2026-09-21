from io import BytesIO
from pathlib import Path
import requests
from bs4 import BeautifulSoup
from pypdf import PdfReader
from docx import Document as Docx
from pptx import Presentation
from openpyxl import load_workbook
def load_pdf(b):return '\n\n'.join(f'[PAGE {i}] {p.extract_text() or ""}' for i,p in enumerate(PdfReader(BytesIO(b)).pages,1))
def load_docx(b):return '\n'.join(p.text for p in Docx(BytesIO(b)).paragraphs if p.text.strip())
def load_pptx(b):
 r=[]
 for i,s in enumerate(Presentation(BytesIO(b)).slides,1):r.append(f'[SLIDE {i}]\n'+'\n'.join(x.text for x in s.shapes if hasattr(x,'text')))
 return '\n\n'.join(r)
def load_xlsx(b):
 w=load_workbook(BytesIO(b),read_only=True,data_only=True);o=[]
 for s in w.worksheets:
  o.append(f'[SHEET {s.title}]')
  o += [' | '.join('' if x is None else str(x) for x in row) for row in s.iter_rows(values_only=True)]
 return '\n'.join(o)
def load_url(url):
 r=requests.get(url,timeout=20,headers={'User-Agent':'EnterpriseAIKnowledgeAssistant'});r.raise_for_status();s=BeautifulSoup(r.text,'html.parser')
 for t in s(['script','style','noscript']):t.decompose()
 return s.get_text('\n',strip=True)
def load_file(name,b):
 e=Path(name).suffix.lower()
 if e=='.pdf':return load_pdf(b)
 if e=='.docx':return load_docx(b)
 if e=='.pptx':return load_pptx(b)
 if e=='.xlsx':return load_xlsx(b)
 if e in {'.txt','.md','.csv'}:return b.decode('utf-8','ignore')
 raise ValueError(f'Unsupported file type: {e}')
